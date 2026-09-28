"""Tests for scripts/human_audit.py (blinded human correctness audit).

Two kinds of fixture:

* the REAL release (data/systems.json, data/coding_frame.csv, data/coded/): the committed draw in
  data/audit/sample.csv must be reproducible from its seed, meet the strata floors and contain no
  double-coded system, and the committed human sheet must carry no model reading;
* a SYNTHETIC root (hand-built systems and frame, the real schema and manual) for the analysis,
  whose expected numbers can be written out by hand.
"""

import csv
import json
import math
import shutil
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

import human_audit as ha

AUDIT = ROOT / "data" / "audit"
HAVE_RELEASE = (ROOT / "data" / "systems.json").exists()


# ------------------------------------------------------------------------------- real release


@pytest.fixture(scope="module")
def release():
    if not HAVE_RELEASE:
        pytest.skip("release not present")
    return ha.load_inputs(ROOT)


def test_sample_reproducible_from_seed(release):
    a = ha.draw_sample(release["systems"], release["frame"], release["excluded"])
    b = ha.draw_sample(release["systems"], release["frame"], release["excluded"])
    assert [r["system_id"] for r in a] == [r["system_id"] for r in b]
    c = ha.draw_sample(release["systems"], release["frame"], release["excluded"], seed="other")
    assert {r["system_id"] for r in a} != {r["system_id"] for r in c}
    if (AUDIT / "sample.csv").exists():
        committed = ha.read_csv(AUDIT / "sample.csv")
        assert [r["system_id"] for r in committed] == [r["system_id"] for r in a]
        assert {r["seed"] for r in committed} == {ha.SEED}


def test_sample_strata_floor_and_allocation(release):
    rows = ha.draw_sample(release["systems"], release["frame"], release["excluded"])
    assert len(rows) == ha.N_SYSTEMS == len({r["system_id"] for r in rows})
    counts = {s: sum(r["stratum"] == s for r in rows) for s in ha.STRATA}
    assert all(counts[s] >= ha.STRATUM_FLOOR for s in ha.STRATA)
    assert counts == {"H": 10, "P": 10, "O": 30}


def test_no_double_coded_or_weightless_system_drawn(release):
    ids = ha.double_coded_ids(ROOT / "data" / "coded")
    assert len(ids) == 247
    rows = ha.draw_sample(release["systems"], release["frame"], release["excluded"])
    drawn = {r["system_id"] for r in rows}
    assert not drawn & ids
    assert not drawn & set(ha.MANUAL_EXAMPLES)
    assert all(float(release["frame"][s]["weight"]) > 0 for s in drawn)


def test_committed_sheet_is_blind():
    if not (AUDIT / "human_sheet.csv").exists():
        pytest.skip("pack not built")
    sheet = ha.read_csv(AUDIT / "human_sheet.csv")
    model = ha.read_csv(AUDIT / "model_answers.csv")
    assert_blind(sheet, model)
    assert len(sheet) == 350


# ------------------------------------------------------------------------------- blinding check


def assert_blind(sheet: list[dict], model: list[dict]) -> None:
    cols = set(sheet[0])
    assert cols == set(ha.SHEET_COLS)
    assert not any(c.startswith("model") or c in ("evidence", "note", "state", "value",
                                                  "confidence", "not_reported")
                   for c in cols)
    # human columns start empty
    assert not any(ha.touched(r) for r in sheet)
    text = {r["cell_id"]: " ".join(r.values()) for r in sheet}
    for m in model:
        for field in ("model_evidence", "model_note"):
            s = (m[field] or "").strip()
            if len(s) >= 25:
                assert s not in text[m["cell_id"]], (m["cell_id"], field)


# ----------------------------------------------------------------------------- synthetic root


N = {"H": 14, "P": 14, "O": 40}
W = {"H": 1.0, "P": 4.0, "O": 40.0}
DOUBLE = ["h00", "h01", "p00", "o00", "o01"]


def _cell(value=None, nr=False, unresolved=False, ev=""):
    c = {"value": value, "not_reported": nr, "confidence": "medium", "coder": "llm-test"}
    if unresolved:
        c["unresolved"] = True
    if ev:
        c["evidence"] = ev
    return c


def _coding(i: int) -> dict:
    """Deterministic model readings: most silence dims NR, C3 coded, E1 alternating."""
    return {
        "network_policy": _cell(nr=True) if i % 4 else _cell("open", ev=f"grep net sys{i} 0 hits"),
        "replayability": _cell(nr=True),
        "filesystem_access": _cell("full", ev=f"root inside container for sys{i} (x.py:1)"),
        "rollback": _cell(nr=True) if i % 3 else _cell("none", ev=f"no undo in registry sys{i}"),
        "state_persistence": _cell(unresolved=True) if i == 5 else _cell(nr=True),
        "multi_agent_topology": _cell("single", ev=f"one agent factory sys{i} (a.py:9)"),
        "self_verification": _cell(["test_execution", "self_critique"],
                                   ev=f"runs pytest then reviews sys{i} (b.py:3)")
        if i % 2 else _cell(nr=True),
        "pinned_version": _cell(f"v1.{i} @ abc{i:04d} (2026-01-01)", ev="tag"),
    }


@pytest.fixture()
def synth(tmp_path: Path) -> Path:
    root = tmp_path / "root"
    (root / "data" / "coded" / "json_pass2").mkdir(parents=True)
    (root / "schema").mkdir()
    (root / "docs").mkdir()
    shutil.copy(ROOT / "schema" / "dimensions.json", root / "schema" / "dimensions.json")
    shutil.copy(ROOT / "docs" / "coding_manual.md", root / "docs" / "coding_manual.md")
    systems, frame = [], []
    i = 0
    for st in ha.STRATA:
        for k in range(N[st]):
            sid = f"{st.lower()}{k:02d}"
            papers = [f"arxiv:{i:04d}"] + (["arxiv:shared"] if k in (2, 3) else [])
            systems.append({"id": sid, "name": sid.upper(), "papers": papers,
                            "urls": {"repo": f"https://github.com/x/{sid}"}, "coding": _coding(i),
                            "notes": "SYSTEM NOTE MUST NOT LEAK"})
            frame.append({"system_id": sid, "stratum": st, "coded": "1", "weight": W[st]})
            i += 1
    # one weight-0 system that must never be drawn
    systems.append({"id": "o-zero", "name": "Zero", "papers": [], "coding": _coding(99)})
    frame.append({"system_id": "o-zero", "stratum": "O", "coded": "1", "weight": 0})
    (root / "data" / "systems.json").write_text(json.dumps(systems), encoding="utf-8")
    ha.write_csv(root / "data" / "coding_frame.csv", ["system_id", "stratum", "coded", "weight"],
                 frame)
    (root / "data" / "coded" / "double_sample.json").write_text(
        json.dumps({"system_ids": DOUBLE[:-1]}), encoding="utf-8")
    (root / "data" / "coded" / "json_pass2" / "o01.json").write_text(
        json.dumps({"system_id": DOUBLE[-1]}), encoding="utf-8")
    return root


def test_allocate_floor_and_proportional():
    assert ha.allocate({"H": 983, "P": 683, "O": 4837}, 50, 10) == {"H": 10, "P": 10, "O": 30}
    # floors do not bind: plain proportional, largest remainder
    assert ha.allocate({"A": 50, "B": 30, "C": 20}, 50, 5) == {"A": 25, "B": 15, "C": 10}
    assert sum(ha.allocate({"A": 1, "B": 1, "C": 1}, 50, 10).values()) == 50
    with pytest.raises(ha.AuditError):
        ha.allocate({"A": 1, "B": 1}, 10, 6)


def test_paper_clusters_union():
    cl = ha.paper_clusters(["a", "b", "c", "d"], {"a": ["p1"], "b": ["p2", "p1"], "c": ["p2"],
                                                  "d": ["p9"]})
    assert cl["a"] == cl["b"] == cl["c"] == "a"
    assert cl["d"] == "d"


def build_synth(root: Path) -> Path:
    out = root / "data" / "audit"
    info = ha.build(root, out)
    assert info["n_cells"] == 7 * info["n_systems"]
    return out


def test_synthetic_build_strata_exclusions_blind(synth: Path):
    out = build_synth(synth)
    sample = ha.read_csv(out / "sample.csv")
    counts = {s: sum(r["stratum"] == s for r in sample) for s in ha.STRATA}
    assert all(v >= ha.STRATUM_FLOOR for v in counts.values())
    assert sum(counts.values()) == ha.N_SYSTEMS
    drawn = {r["system_id"] for r in sample}
    assert not drawn & set(DOUBLE)
    assert "o-zero" not in drawn
    sheet = ha.read_csv(out / "human_sheet.csv")
    model = ha.read_csv(out / "model_answers.csv")
    assert_blind(sheet, model)
    assert not any("SYSTEM NOTE" in " ".join(r.values()) for r in sheet)
    # the model's values for audited dims never appear as a sheet column value
    assert all(r["human_value"] == "" and r["human_state"] == "" for r in sheet)
    assert [r["cell_id"] for r in sheet] == [r["cell_id"] for r in model]
    # glosses and rules come from the manual
    g3 = next(r for r in sheet if r["dim_id"] == "G3")
    assert "open: no restriction" in g3["value_glosses"]
    assert g3["decision_rule"].startswith("Definition: ")
    # stable: a rebuild gives the same draw
    ha.build(synth, out)
    assert ha.read_csv(out / "sample.csv") == sample


def test_build_refuses_to_overwrite_human_work(synth: Path):
    out = build_synth(synth)
    sheet = ha.read_csv(out / "human_sheet.csv")
    sheet[0]["human_state"] = "coded"
    ha.write_csv(out / "human_sheet.csv", ha.SHEET_COLS, sheet)
    with pytest.raises(ha.AuditError, match="refusing to overwrite"):
        ha.build(synth, out)


def fill_as_model(out: Path) -> list[dict]:
    """A human sheet that agrees with the model everywhere (unresolved model cells -> NR)."""
    sheet = ha.read_csv(out / "human_sheet.csv")
    model = {r["cell_id"]: r for r in ha.read_csv(out / "model_answers.csv")}
    for r in sheet:
        m = model[r["cell_id"]]
        st = m["model_state"] if m["model_state"] != "unresolved" else "not_reported"
        r["human_state"] = st
        r["human_value"] = m["model_value"] if st == "coded" else ""
        r["human_evidence_quote"] = "quote" if st == "coded" else ""
        r["human_locator"] = "x.py:1@abc"
        r["human_confidence"] = "medium"
        r["human_minutes"] = "4"
        r["coder_id"] = "c1"
    return sheet


def test_analysis_refuses_empty_sheet_and_writes_no_tex(synth: Path):
    out = build_synth(synth)
    tex = synth / "paper" / "tables" / "human_audit.tex"
    with pytest.raises(ha.AuditError, match="empty"):
        ha.run_analysis(synth, out, out / "human_sheet.csv", tex)
    with pytest.raises(ha.AuditError, match="empty"):
        ha.run_analysis(synth, out, out / "human_sheet.csv", tex, allow_partial=True)
    assert not tex.exists()
    assert not (out / "results.json").exists()
    assert ha.main(["--analyse", "--root", str(synth)]) == 2
    assert not tex.exists()


def test_analysis_partial_sheet(synth: Path):
    out = build_synth(synth)
    sheet = fill_as_model(out)
    for r in sheet[70:]:
        for c in ha.HUMAN_COLS:
            r[c] = ""
    ha.write_csv(out / "human_sheet.csv", ha.SHEET_COLS, sheet)
    tex = synth / "t.tex"
    with pytest.raises(ha.AuditError, match="--allow-partial"):
        ha.run_analysis(synth, out, out / "human_sheet.csv", tex)
    res = ha.run_analysis(synth, out, out / "human_sheet.csv", tex, allow_partial=True, draws=50)
    assert res["status"] == "partial" and res["n_filled"] == 70
    assert not tex.exists()


def test_analysis_rejects_invalid_rows(synth: Path):
    out = build_synth(synth)
    sheet = fill_as_model(out)
    g3 = next(r for r in sheet if r["dim_id"] == "G3")
    g3.update(human_state="coded", human_value="open|none", human_evidence_quote="q")
    e1 = next(r for r in sheet if r["dim_id"] == "E1")
    e1.update(human_state="coded", human_value="none|test_execution", human_evidence_quote="q")
    c3 = next(r for r in sheet if r["dim_id"] == "C3")
    c3.update(human_state="not_reported", human_value="single")
    ha.write_csv(out / "human_sheet.csv", ha.SHEET_COLS, sheet)
    with pytest.raises(ha.AuditError) as e:
        ha.run_analysis(synth, out, out / "human_sheet.csv", None)
    msg = str(e.value)
    assert "single-valued" in msg and "rule 4" in msg and "must not carry a value" in msg


def test_analysis_on_synthetic_filled_sheet(synth: Path):
    out = build_synth(synth)
    sheet = fill_as_model(out)
    model = {r["cell_id"]: r for r in ha.read_csv(out / "model_answers.csv")}
    # Inject known disagreements on distinct cells.
    rb = [r for r in sheet if r["dim_id"] == "E3" and model[r["cell_id"]]["model_state"]
          == "not_reported"][:3]
    for r in rb:  # rule-5 split: human finds the absence value
        r.update(human_state="coded", human_value="none", human_evidence_quote="registry, no undo")
    c3 = [r for r in sheet if r["dim_id"] == "C3"][:2]
    for r in c3:  # value mismatch
        r.update(human_value="pipeline", human_evidence_quote="three agents in sequence")
    e1 = [r for r in sheet if r["dim_id"] == "E1" and model[r["cell_id"]]["model_state"]
          == "coded"][:1]
    for r in e1:  # partial overlap on a multi-valued dimension
        r.update(human_value="test_execution")
    h2 = [r for r in sheet if r["dim_id"] == "H2"][:1]
    for r in h2:  # human cannot settle it: excluded from the accuracy denominator
        r.update(human_state="unresolved", human_value="", human_evidence_quote="",
                 human_note="ran out of time")
    ha.write_csv(out / "human_sheet.csv", ha.SHEET_COLS, sheet)
    tex = synth / "paper" / "tables" / "human_audit.tex"
    res = ha.run_analysis(synth, out, out / "human_sheet.csv", tex, draws=200)

    assert res["status"] == "complete" and res["n_filled"] == 350
    # the model's single unresolved cell (system index 5, D3) is wrong against a human NR
    n_model_unres = sum(m["model_state"] == "unresolved" for m in model.values())
    n_wrong = 3 + 2 + 1 + n_model_unres
    p = res["pooled"]
    assert p["n_ref"] == 349
    assert p["n_correct"] == 349 - n_wrong
    assert math.isclose(p["accuracy"], (349 - n_wrong) / 349)
    b = p["paper_cluster_bootstrap"]
    assert b["accuracy_ci_lo"] <= p["accuracy"] <= b["accuracy_ci_hi"]
    # systems k = 2, 3 of every stratum share one paper and must form a single cluster
    drawn = {r["system_id"] for r in ha.read_csv(out / "sample.csv")}
    sharing = drawn & {f"{s}{k:02d}" for s in "hpo" for k in (2, 3)}
    assert b["n_clusters"] == 50 - len(sharing) + (1 if sharing else 0)

    e3 = res["per_dimension"]["rollback"]
    assert e3["n_ref"] == 50 and e3["n_correct"] == 47
    acc, lo, _hi = ha.wilson(47, 50)
    assert math.isclose(e3["accuracy"], acc) and math.isclose(e3["accuracy_ci_lo"], lo)
    assert e3["disagreement_types"] == {"nr_vs_absence": 3}
    assert e3["state_confusion_model_by_human"]["not_reported"]["coded"] == 3
    assert e3["silence_recall"] == 1.0 and e3["silence_precision"] < 1.0
    c3r = res["per_dimension"]["multi_agent_topology"]
    assert c3r["disagreement_types"] == {"value_mismatch_disjoint": 2}
    assert c3r["n_both_coded"] == 50 and math.isclose(c3r["value_agreement_given_coded"], 0.96)
    e1r = res["per_dimension"]["self_verification"]
    assert e1r["disagreement_types"] == {"value_mismatch_partial": 1}
    assert "per_value_kappa" in e1r
    h2r = res["per_dimension"]["replayability"]
    assert h2r["n_human_unresolved"] == 1 and h2r["n_ref"] == 49
    assert h2r["disagreement_types"] == {"unresolved": 1}
    g3r = res["per_dimension"]["network_policy"]
    assert g3r["accuracy"] == 1.0 and g3r["kappa"] == 1.0 and g3r["ac1"] == 1.0

    # disagreement table carries both evidence strings
    d = next(x for x in res["disagreements"] if x["dim_key"] == "multi_agent_topology")
    assert d["model_evidence"].startswith("one agent factory")
    assert d["human_evidence"] == "three agents in sequence"
    assert len(res["disagreements"]) == n_wrong + 1

    # outputs
    j = json.loads((out / "results.json").read_text(encoding="utf-8"))
    assert j["pooled"]["n_correct"] == p["n_correct"]
    md = (out / "results.md").read_text(encoding="utf-8")
    assert "three agents in sequence" in md and "nr_vs_absence" in md
    t = tex.read_text(encoding="utf-8")
    assert "\\label{tab:human-audit}" in t and "E3 rollback & 50 & 0.94" in t

    # weighted accuracy is a proper weighted mean of the correct flags
    assert 0 <= p["weighted_accuracy"] <= 1


def test_analysis_reads_xlsx(synth: Path):
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook

    out = build_synth(synth)
    sheet = fill_as_model(out)
    wb = load_workbook(out / "human_sheet.xlsx")
    ws = wb["cells"]
    header = [c.value for c in ws[1]]
    for i, r in enumerate(sheet, 2):
        for c in ha.HUMAN_COLS:
            ws.cell(row=i, column=header.index(c) + 1, value=r[c] or None)
    wb.save(out / "human_sheet.xlsx")
    assert ha.pick_sheet(out, None).suffix == ".xlsx"
    res = ha.run_analysis(synth, out, out / "human_sheet.xlsx", None, draws=20)
    assert res["status"] == "complete"


def test_wilson_reference_values():
    _p, lo, hi = ha.wilson(45, 50)
    assert (round(lo, 3), round(hi, 3)) == (0.786, 0.957)
    assert all(math.isnan(x) for x in ha.wilson(0, 0))


def test_csv_roundtrip_keeps_columns(tmp_path: Path):
    ha.write_csv(tmp_path / "x.csv", ["a", "b"], [{"a": "1|2", "b": ""}])
    with (tmp_path / "x.csv").open(encoding="utf-8") as fh:
        assert list(csv.DictReader(fh)) == [{"a": "1|2", "b": ""}]
