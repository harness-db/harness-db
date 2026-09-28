#!/usr/bin/env python
"""Blinded human correctness audit of HARNESS-DB cells (protocol: docs/human_audit_protocol.md).

Model-model reliability (``docs/coding_reliability.md``) says the coding procedure is reproducible;
it cannot say it is right. This script builds the pack for a blinded human reading of 350 cells
(50 systems x 7 dimensions) and, once a human has filled it, scores the released model reading
against the human one.

Build the pack (sample, blinded sheet, model answers kept apart):

    python scripts/human_audit.py --build

Analyse a filled sheet (refuses an empty one; writes the LaTeX fragment only for a complete one):

    python scripts/human_audit.py --analyse [--sheet data/audit/human_sheet.xlsx] [--allow-partial]

Outputs live in ``data/audit/``; the LaTeX fragment goes to ``paper/tables/human_audit.tex``.
"""
from __future__ import annotations

import argparse
import csv
import json
import math
import random
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import kappa

ROOT = Path(__file__).resolve().parents[1]

SEED = "human-audit-2026-09-28"
N_SYSTEMS = 50
STRATUM_FLOOR = 10
STRATA = ("H", "P", "O")
AUDIT_DIMS = ("G3", "H2", "G2", "E3", "D3", "C3", "E1")
# The dimension's "absence-type" value: used only to type disagreements (protocol 4.4).
ABSENCE_VALUE = {
    "network_policy": "open",
    "replayability": "none",
    "filesystem_access": "full",
    "rollback": "none",
    "state_persistence": "none",
    "multi_agent_topology": "single",
    "self_verification": "none",
}
# Printed as calibration examples in docs/coding_manual.md, so never drawn.
MANUAL_EXAMPLES = ("swe-agent", "openhands")
VERSION_CUTOFF = "2026-08-31"
BOOTSTRAP_DRAWS = 2000
BOOTSTRAP_SEED = 20260928

STATE_CODED = "coded"
STATE_NR = "not_reported"
STATE_UNRESOLVED = "unresolved"
STATES = (STATE_CODED, STATE_NR, STATE_UNRESOLVED)
CONFIDENCES = ("high", "medium", "low")

# Fallback glosses for values the manual lists without one.
FALLBACK_GLOSS = {
    "none": "the mechanism is absent: you opened the place it would be declared and it is not "
            "there (rule 5b)",
    "read_only": "the agent may read files but not write them",
}

THREE_STATE_RULE = (
    "coded = you can point at evidence for a value; for an absence value (none, open, ...) you must "
    "have opened the place where the feature would be declared if it existed (config schema, CLI "
    "flags, tool registry, the run loop, a documented feature list) and found it missing (rule 5b). "
    "not_reported = the pinned sources contain no such place, or you could not find one (rule 5c); "
    "prose that does not mention the feature is NOT evidence of absence. unresolved = you cannot "
    "settle it within the 8-minute time box; say why in human_note."
)

SHEET_INFO_COLS = [
    "row", "cell_id", "system_id", "system_name", "pinned_version", "repo_url", "paper_ids",
    "paper_refs", "dim_id", "dim_key", "dim_name", "multi_valued", "permitted_values",
    "value_glosses", "decision_rule", "three_state_rule",
]
HUMAN_COLS = [
    "human_state", "human_value", "human_evidence_quote", "human_locator", "human_confidence",
    "human_minutes", "human_note", "coder_id",
]
SHEET_COLS = SHEET_INFO_COLS + HUMAN_COLS
MODEL_COLS = [
    "cell_id", "system_id", "dim_id", "dim_key", "model_state", "model_value", "model_evidence",
    "model_confidence", "model_note", "model_coder",
]
SAMPLE_COLS = [
    "seed", "sheet_order", "stratum", "draw_order", "system_id", "system_name", "frame_weight",
    "eligible_in_stratum", "sampled_in_stratum", "audit_weight", "paper_cluster", "paper_ids",
]

csv.field_size_limit(10 ** 8)


class AuditError(Exception):
    """The inputs or the filled sheet cannot be used; the message says why."""


# --------------------------------------------------------------------------------------- loading


def read_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def write_csv(path: Path, cols: list[str], rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=cols, lineterminator="\n")
        w.writeheader()
        for r in rows:
            w.writerow({c: r.get(c, "") for c in cols})


def audit_dims(dims_schema: dict) -> list[dict]:
    by_id = {d["id"]: d for d in dims_schema["dimensions"]}
    return [by_id[i] for i in AUDIT_DIMS]


def double_coded_ids(coded_dir: Path) -> set[str]:
    """Union of the reliability sample ids: double_sample.json and the json_pass2 codings."""
    ids: set[str] = set()
    ds = coded_dir / "double_sample.json"
    if ds.exists():
        ids.update(read_json(ds).get("system_ids", []))
    p2 = coded_dir / "json_pass2"
    if p2.is_dir():
        for f in sorted(p2.glob("*.json")):
            try:
                d = read_json(f)
            except (json.JSONDecodeError, UnicodeDecodeError):
                d = {}
            ids.add(d.get("system_id") or f.stem)
    return ids


# ------------------------------------------------------------------------------- cell semantics


def cell_state(cell: dict | None) -> str:
    """Same three-state mapping as scripts/analyse_descriptives.py ``cell_state``."""
    cell = cell or {}
    if cell.get("unresolved"):
        return STATE_UNRESOLVED
    if cell.get("not_reported"):
        return STATE_NR
    if cell.get("value") is None or cell.get("value") == []:
        return STATE_UNRESOLVED
    return STATE_CODED


def value_set(value) -> tuple[str, ...]:
    if value is None:
        return ()
    if isinstance(value, list):
        vals = value
    else:
        vals = str(value).split("|")
    return tuple(sorted({str(v).strip() for v in vals if str(v).strip()}))


def label(state: str, values: tuple[str, ...]) -> str:
    """The comparison label: NR, UNRESOLVED or the sorted pipe-joined value set (as kappa.norm)."""
    if state == STATE_NR:
        return "NR"
    if state == STATE_UNRESOLVED:
        return "UNRESOLVED"
    return "|".join(values)


# ------------------------------------------------------------------------------------- sampling


def eligible_by_stratum(systems: list[dict], frame: dict[str, dict], excluded: set[str]
                        ) -> tuple[dict[str, list[str]], dict[str, float]]:
    """Weight-bearing released systems per stratum (minus exclusions), and each stratum's weighted
    frame total over ALL weight-bearing released systems (the allocation base)."""
    elig: dict[str, list[str]] = {s: [] for s in STRATA}
    pop: dict[str, float] = {s: 0.0 for s in STRATA}
    for s in systems:
        fr = frame.get(s["id"])
        if not fr:
            continue
        w = float(fr.get("weight") or 0.0)
        if w <= 0 or str(fr.get("coded", "1")) != "1" or fr.get("stratum") not in STRATA:
            continue
        pop[fr["stratum"]] += w
        if s["id"] not in excluded:
            elig[fr["stratum"]].append(s["id"])
    return {k: sorted(v) for k, v in elig.items()}, pop


def allocate(pop: dict[str, float], n: int, floor: int) -> dict[str, int]:
    """Proportional allocation with a per-stratum floor, largest-remainder rounding.

    Strata whose proportional quota falls below ``floor`` are fixed at ``floor`` and the rest of
    ``n`` is re-allocated proportionally among the others, until no quota is below the floor.
    """
    strata = [s for s in pop if pop[s] > 0]
    if floor * len(strata) > n:
        raise AuditError(f"floor {floor} x {len(strata)} strata exceeds n = {n}")
    fixed: dict[str, int] = {}
    while True:
        free = [s for s in strata if s not in fixed]
        rem = n - sum(fixed.values())
        tot = sum(pop[s] for s in free)
        quota = {s: rem * pop[s] / tot for s in free} if tot else {}
        low = [s for s in free if quota[s] < floor]
        if not low:
            break
        for s in low:
            fixed[s] = floor
    alloc = dict(fixed)
    base = {s: math.floor(q) for s, q in quota.items()}
    left = n - sum(alloc.values()) - sum(base.values())
    order = sorted(quota, key=lambda s: (-(quota[s] - base[s]), s))
    for s in order[:left]:
        base[s] += 1
    alloc.update(base)
    return {s: alloc.get(s, 0) for s in pop}


def paper_clusters(system_ids: list[str], papers: dict[str, list[str]]) -> dict[str, str]:
    """Union-find over the sample: systems sharing any paper id are one cluster.

    The cluster label is the lexicographically smallest system id in it.
    """
    parent = {s: s for s in system_ids}

    def find(x: str) -> str:
        while parent[x] != x:
            parent[x] = parent[parent[x]]
            x = parent[x]
        return x

    owner: dict[str, str] = {}
    for sid in system_ids:
        for p in papers.get(sid, []):
            if p in owner:
                a, b = find(owner[p]), find(sid)
                if a != b:
                    parent[max(a, b)] = min(a, b)
            else:
                owner[p] = sid
    return {s: find(s) for s in system_ids}


def draw_sample(systems: list[dict], frame: dict[str, dict], excluded: set[str],
                seed: str = SEED, n: int = N_SYSTEMS, floor: int = STRATUM_FLOOR) -> list[dict]:
    """The reproducible stratified draw; one dict per sampled system, in sheet order."""
    by_id = {s["id"]: s for s in systems}
    elig, pop = eligible_by_stratum(systems, frame, excluded)
    alloc = allocate(pop, n, floor)
    rows = []
    for st in STRATA:
        k = alloc.get(st, 0)
        if k > len(elig[st]):
            raise AuditError(f"stratum {st}: {k} to draw but only {len(elig[st])} eligible")
        picked = random.Random(f"{seed}:{st}").sample(elig[st], k)
        for i, sid in enumerate(picked, 1):
            fw = float(frame[sid]["weight"])
            rows.append({
                "seed": seed, "stratum": st, "draw_order": i, "system_id": sid,
                "system_name": by_id[sid].get("name", sid), "frame_weight": fw,
                "eligible_in_stratum": len(elig[st]), "sampled_in_stratum": k,
                "audit_weight": round(fw * len(elig[st]) / k, 6),
                "paper_ids": "|".join(by_id[sid].get("papers", [])),
            })
    # Interleave strata in the sheet so fatigue and learning are not confounded with stratum.
    order = [r["system_id"] for r in rows]
    random.Random(f"{seed}:order").shuffle(order)
    pos = {sid: i for i, sid in enumerate(order, 1)}
    clusters = paper_clusters(order, {s: by_id[s].get("papers", []) for s in order})
    for r in rows:
        r["sheet_order"] = pos[r["system_id"]]
        r["paper_cluster"] = clusters[r["system_id"]]
    return sorted(rows, key=lambda r: r["sheet_order"])


# ------------------------------------------------------------------------------ manual parsing


def manual_section(manual: str, dim: dict) -> str:
    m = re.search(rf"^#### {re.escape(dim['id'])} {re.escape(dim['key'])}\b.*$", manual, re.MULTILINE)
    if not m:
        raise AuditError(f"coding manual has no section for {dim['id']} {dim['key']}")
    rest = manual[m.end():]
    end = re.search(r"^#{2,4} ", rest, re.MULTILINE)
    return rest[: end.start()] if end else rest


def manual_glosses(manual: str, dim: dict) -> dict[str, str]:
    """Value -> one-line gloss from the dimension's "Values:" list in the manual."""
    sec = manual_section(manual, dim)
    out: dict[str, str] = {}
    for m in re.finditer(r"^- `([a-z_0-9]+)`(?::\s*(.*))?$", sec, re.MULTILINE):
        out[m.group(1)] = (m.group(2) or "").strip()
    glosses = {}
    for v in dim["values"]:
        g = out.get(v) or FALLBACK_GLOSS.get(v, "")
        glosses[v] = g
    return glosses


def manual_rule(manual: str, dim: dict) -> str:
    sec = manual_section(manual, dim)
    m = re.search(r"^Decision rule:(.*?)(?:\n\s*\n|\Z)", sec, re.MULTILINE | re.DOTALL)
    head = re.search(r"^Definition:(.*?)(?:\n\s*\n|\Z)", sec, re.MULTILINE | re.DOTALL)
    parts = []
    if head:
        parts.append("Definition: " + " ".join(head.group(1).split()))
    if m:
        parts.append("Decision rule: " + " ".join(m.group(1).split()))
    return " ".join(parts)


# ---------------------------------------------------------------------------------- sheet build


def pinned_text(system: dict) -> str:
    cell = system.get("coding", {}).get("pinned_version") or {}
    if cell_state(cell) == STATE_CODED:
        return str(cell["value"])
    label_ = system.get("version_label") or "-"
    return (f"no pin recorded (version label: {label_}); pin per protocol 4.5 - latest tag on or "
            f"before {VERSION_CUTOFF}, else latest default-branch commit on or before it - and "
            "record the commit you read in human_locator")


def paper_refs(ids: list[str], papers: dict[str, dict]) -> str:
    out = []
    for p in ids:
        rec = papers.get(p) or {}
        title = (rec.get("title") or "").strip()
        url = (rec.get("url") or "").strip()
        if not url and rec.get("arxiv_id"):
            url = f"https://arxiv.org/abs/{rec['arxiv_id']}"
        out.append(" ".join(x for x in (p, f"'{title}'" if title else "", f"<{url}>" if url else "")
                            if x))
    return " ; ".join(out)


def build_rows(sample: list[dict], systems: list[dict], dims: list[dict], manual: str,
               papers: dict[str, dict]) -> tuple[list[dict], list[dict]]:
    """(human sheet rows, model answer rows) for the sample, in the same cell order."""
    by_id = {s["id"]: s for s in systems}
    meta = {d["key"]: (manual_glosses(manual, d), manual_rule(manual, d)) for d in dims}
    sheet, model = [], []
    n = 0
    for srow in sample:
        sys_ = by_id[srow["system_id"]]
        for d in dims:
            n += 1
            glosses, rule = meta[d["key"]]
            cid = f"{sys_['id']}::{d['id']}"
            sheet.append({
                "row": n, "cell_id": cid, "system_id": sys_["id"],
                "system_name": sys_.get("name", ""), "pinned_version": pinned_text(sys_),
                "repo_url": (sys_.get("urls") or {}).get("repo", ""),
                "paper_ids": "|".join(sys_.get("papers", [])),
                "paper_refs": paper_refs(sys_.get("papers", []), papers),
                "dim_id": d["id"], "dim_key": d["key"], "dim_name": d["name"],
                "multi_valued": "yes (pipe-join values)" if d.get("multi") else "no",
                "permitted_values": "|".join(d["values"]),
                "value_glosses": " ; ".join(f"{v}: {g}" if g else v for v, g in glosses.items()),
                "decision_rule": rule, "three_state_rule": THREE_STATE_RULE,
                **{c: "" for c in HUMAN_COLS},
            })
            cell = sys_.get("coding", {}).get(d["key"]) or {}
            st = cell_state(cell)
            model.append({
                "cell_id": cid, "system_id": sys_["id"], "dim_id": d["id"], "dim_key": d["key"],
                "model_state": st,
                "model_value": "|".join(value_set(cell.get("value"))) if st == STATE_CODED else "",
                "model_evidence": cell.get("evidence", "") or "",
                "model_confidence": cell.get("confidence", "") or "",
                "model_note": cell.get("note", "") or "", "model_coder": cell.get("coder", "") or "",
            })
    return sheet, model


def write_xlsx(path: Path, rows: list[dict], dims: list[dict], readme: str) -> bool:
    """The same sheet as an .xlsx with dropdowns; returns False if openpyxl is missing."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        return False
    wb = Workbook()
    ws = wb.active
    ws.title = "cells"
    ws.append(SHEET_COLS)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([r.get(c, "") for c in SHEET_COLS])
    ws.freeze_panes = "E2"
    col = {c: get_column_letter(i) for i, c in enumerate(SHEET_COLS, 1)}
    widths = {"value_glosses": 60, "decision_rule": 60, "three_state_rule": 40, "paper_refs": 40,
              "pinned_version": 30, "repo_url": 30, "human_evidence_quote": 40,
              "human_locator": 30, "human_note": 30}
    for c, letter in col.items():
        ws.column_dimensions[letter].width = widths.get(c, 16)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    last = len(rows) + 1

    def dv_list(options: list[str], strict: bool, prompt: str) -> DataValidation:
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"',
                            allow_blank=True, showErrorMessage=strict)
        dv.promptTitle, dv.prompt, dv.showInputMessage = "HARNESS-DB audit", prompt, True
        ws.add_data_validation(dv)
        return dv

    dv_list(list(STATES), True, "coded / not_reported / unresolved (rule 5)").add(
        f"{col['human_state']}2:{col['human_state']}{last}")
    dv_list(list(CONFIDENCES), True, "general rule 3").add(
        f"{col['human_confidence']}2:{col['human_confidence']}{last}")
    for d in dims:
        multi = bool(d.get("multi"))
        prompt = ("pick one value, or type several joined with |" if multi
                  else "one value; empty unless human_state = coded")
        dv = dv_list(list(d["values"]), not multi, prompt)
        for i, r in enumerate(rows, 2):
            if r["dim_key"] == d["key"]:
                dv.add(f"{col['human_value']}{i}")
    ws2 = wb.create_sheet("procedure")
    for line in readme.splitlines():
        ws2.append([line])
    ws2.column_dimensions["A"].width = 120
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return True


def read_sheet(path: Path) -> list[dict[str, str]]:
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb["cells"] if "cells" in wb.sheetnames else wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = [str(h) if h is not None else "" for h in next(it)]
        out = []
        for vals in it:
            if vals is None or all(v is None for v in vals):
                continue
            out.append({h: ("" if v is None else str(v)) for h, v in zip(header, vals)})
        wb.close()
        return out
    return read_csv(path)


def filled(row: dict) -> bool:
    return bool((row.get("human_state") or "").strip())


def touched(row: dict) -> bool:
    return any((row.get(c) or "").strip() for c in HUMAN_COLS)


def load_inputs(root: Path) -> dict:
    systems = read_json(root / "data" / "systems.json")
    frame = {r["system_id"]: r for r in read_csv(root / "data" / "coding_frame.csv")}
    dims = audit_dims(read_json(root / "schema" / "dimensions.json"))
    manual = (root / "docs" / "coding_manual.md").read_text(encoding="utf-8")
    papers_path = root / "data" / "papers.csv"
    papers = {r["id"]: r for r in read_csv(papers_path)} if papers_path.exists() else {}
    excluded = double_coded_ids(root / "data" / "coded") | set(MANUAL_EXAMPLES)
    return {"systems": systems, "frame": frame, "dims": dims, "manual": manual, "papers": papers,
            "excluded": excluded}


def build(root: Path, out_dir: Path, force: bool = False) -> dict:
    """Write sample.csv, human_sheet.csv(/.xlsx) and model_answers.csv into ``out_dir``."""
    for name in ("human_sheet.csv", "human_sheet.xlsx"):
        p = out_dir / name
        if p.exists() and not force and any(touched(r) for r in read_sheet(p)):
            raise AuditError(f"{p} already carries human entries; refusing to overwrite "
                             "(pass --force only if you mean to discard them)")
    inp = load_inputs(root)
    sample = draw_sample(inp["systems"], inp["frame"], inp["excluded"])
    sheet, model = build_rows(sample, inp["systems"], inp["dims"], inp["manual"], inp["papers"])
    write_csv(out_dir / "sample.csv", SAMPLE_COLS, sample)
    write_csv(out_dir / "human_sheet.csv", SHEET_COLS, sheet)
    write_csv(out_dir / "model_answers.csv", MODEL_COLS, model)
    readme_p = out_dir / "README.md"
    readme = readme_p.read_text(encoding="utf-8") if readme_p.exists() else ""
    xlsx = write_xlsx(out_dir / "human_sheet.xlsx", sheet, inp["dims"], readme)
    strata = Counter(r["stratum"] for r in sample)
    return {"n_systems": len(sample), "n_cells": len(sheet), "strata": dict(strata),
            "eligible": {r["stratum"]: r["eligible_in_stratum"] for r in sample},
            "n_excluded_double_coded": len(inp["excluded"] - set(MANUAL_EXAMPLES)),
            "xlsx": xlsx}


# ----------------------------------------------------------------------------------- validation


def validate(sheet: list[dict], model: list[dict], dims: list[dict],
             allow_partial: bool) -> list[dict]:
    """Return the filled, validated human rows; raise AuditError with every problem found."""
    ids_sheet = [r.get("cell_id", "") for r in sheet]
    ids_model = {r["cell_id"] for r in model}
    errors = []
    if set(ids_sheet) != ids_model or len(ids_sheet) != len(ids_model):
        errors.append(f"sheet cell ids do not match model_answers.csv ({len(ids_sheet)} rows vs "
                      f"{len(ids_model)} cells); was the sheet rebuilt or edited structurally?")
    n_filled = sum(filled(r) for r in sheet)
    if n_filled == 0:
        raise AuditError("the human sheet is empty: no row has human_state filled. Nothing to "
                         "analyse, and no result may be reported.")
    if n_filled < len(sheet) and not allow_partial:
        raise AuditError(f"only {n_filled} of {len(sheet)} rows are filled; pass --allow-partial "
                         "to analyse a partial sheet (JSON/Markdown only, never the LaTeX table)")
    spec = {d["key"]: d for d in dims}
    out = []
    for r in sheet:
        cid = r.get("cell_id", "?")
        if not filled(r):
            if touched(r):
                errors.append(f"{cid}: human fields filled but human_state is empty")
            continue
        st = r["human_state"].strip().lower()
        d = spec.get(r.get("dim_key", ""))
        if d is None:
            errors.append(f"{cid}: unknown dimension {r.get('dim_key')!r}")
            continue
        vals = value_set(r.get("human_value", ""))
        if st not in STATES:
            errors.append(f"{cid}: human_state {r['human_state']!r} is not one of {STATES}")
            continue
        if st == STATE_CODED:
            if not vals:
                errors.append(f"{cid}: coded but human_value is empty")
            bad = [v for v in vals if v not in d["values"]]
            if bad:
                errors.append(f"{cid}: value(s) {bad} not permitted for {d['key']}")
            if len(vals) > 1 and not d.get("multi"):
                errors.append(f"{cid}: {d['key']} is single-valued but got {vals}")
            if "none" in vals and len(vals) > 1:
                errors.append(f"{cid}: 'none' listed alongside other values (rule 4)")
            if not (r.get("human_evidence_quote") or "").strip():
                errors.append(f"{cid}: coded without an evidence quote (rule 2)")
        elif vals:
            errors.append(f"{cid}: {st} must not carry a value (got {vals})")
        conf = (r.get("human_confidence") or "").strip().lower()
        if conf and conf not in CONFIDENCES:
            errors.append(f"{cid}: human_confidence {conf!r} is not one of {CONFIDENCES}")
        mins = (r.get("human_minutes") or "").strip()
        if mins:
            try:
                float(mins)
            except ValueError:
                errors.append(f"{cid}: human_minutes {mins!r} is not a number")
        out.append({**r, "human_state": st, "_values": vals})
    if errors:
        raise AuditError("the sheet failed validation:\n  " + "\n  ".join(errors))
    return out


# ------------------------------------------------------------------------------------- analysis


def wilson(k: int, n: int, z: float = 1.959964) -> tuple[float, float, float]:
    """(proportion, lower, upper) Wilson score interval; NaNs when n = 0."""
    if n == 0:
        return float("nan"), float("nan"), float("nan")
    p = k / n
    den = 1 + z * z / n
    centre = (p + z * z / (2 * n)) / den
    half = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / den
    return p, max(0.0, centre - half), min(1.0, centre + half)


def disagreement_type(m_state: str, m_vals: tuple, h_state: str, h_vals: tuple,
                      absence: str) -> str:
    if STATE_UNRESOLVED in (m_state, h_state):
        return "unresolved"
    if STATE_NR in (m_state, h_state):
        other = h_vals if m_state == STATE_NR else m_vals
        return "nr_vs_absence" if other == (absence,) else "nr_vs_feature"
    return "value_mismatch_partial" if set(m_vals) & set(h_vals) else "value_mismatch_disjoint"


def _clean(x):
    if isinstance(x, float) and not math.isfinite(x):
        return None
    if isinstance(x, dict):
        return {k: _clean(v) for k, v in x.items()}
    if isinstance(x, (list, tuple)):
        return [_clean(v) for v in x]
    return x


def pooled_accuracy(cells: list[dict]) -> tuple[float, float]:
    ref = [c for c in cells if c["h_state"] != STATE_UNRESOLVED]
    if not ref:
        return float("nan"), float("nan")
    acc = sum(c["correct"] for c in ref) / len(ref)
    wsum = sum(c["weight"] for c in ref)
    wacc = sum(c["weight"] * c["correct"] for c in ref) / wsum if wsum else float("nan")
    return acc, wacc


def cluster_bootstrap_accuracy(cells: list[dict], draws: int, seed: int) -> dict:
    by_cluster: dict[str, list[dict]] = defaultdict(list)
    for c in cells:
        by_cluster[c["cluster"]].append(c)
    keys = sorted(by_cluster)
    rng = random.Random(seed)
    acc_d, wacc_d = [], []
    for _ in range(draws):
        pick = [keys[rng.randrange(len(keys))] for _ in keys]
        a, w = pooled_accuracy([c for k in pick for c in by_cluster[k]])
        acc_d.append(a)
        wacc_d.append(w)
    return {
        "draws": draws, "seed": seed, "n_clusters": len(keys),
        "accuracy_ci_lo": kappa._percentile(acc_d, 0.025),
        "accuracy_ci_hi": kappa._percentile(acc_d, 0.975),
        "weighted_accuracy_ci_lo": kappa._percentile(wacc_d, 0.025),
        "weighted_accuracy_ci_hi": kappa._percentile(wacc_d, 0.975),
    }


def analyse(human: list[dict], model: list[dict], sample: list[dict], dims: list[dict],
            n_sheet_rows: int, draws: int = BOOTSTRAP_DRAWS, seed: int = BOOTSTRAP_SEED) -> dict:
    mby = {r["cell_id"]: r for r in model}
    sby = {r["system_id"]: r for r in sample}
    cells = []
    for h in human:
        m = mby[h["cell_id"]]
        s = sby[h["system_id"]]
        m_vals = value_set(m["model_value"]) if m["model_state"] == STATE_CODED else ()
        h_vals = h["_values"] if h["human_state"] == STATE_CODED else ()
        ml, hl = label(m["model_state"], m_vals), label(h["human_state"], h_vals)
        cells.append({
            "cell_id": h["cell_id"], "system_id": h["system_id"], "dim_key": h["dim_key"],
            "stratum": s["stratum"], "weight": float(s["audit_weight"]),
            "cluster": s["paper_cluster"],
            "m_state": m["model_state"], "h_state": h["human_state"],
            "m_vals": m_vals, "h_vals": h_vals, "m_label": ml, "h_label": hl,
            "correct": ml == hl, "model": m, "human": h,
        })
    per_dim = {}
    disagreements = []
    for d in dims:
        dc = [c for c in cells if c["dim_key"] == d["key"]]
        ref = [c for c in dc if c["h_state"] != STATE_UNRESOLVED]
        k_ok = sum(c["correct"] for c in ref)
        acc, lo, hi = wilson(k_ok, len(ref))
        confusion = {ms: {hs: sum(1 for c in dc if c["m_state"] == ms and c["h_state"] == hs)
                          for hs in STATES} for ms in STATES}
        st_k, st_po, _ = kappa.cohen_kappa([c["m_state"] for c in dc], [c["h_state"] for c in dc])
        both = [c for c in dc if c["m_state"] == c["h_state"] == STATE_CODED]
        v_ok = sum(c["m_vals"] == c["h_vals"] for c in both)
        v_p, v_lo, v_hi = wilson(v_ok, len(both))
        ml, hl = [c["m_label"] for c in dc], [c["h_label"] for c in dc]
        k_lab, po_lab, _ = kappa.cohen_kappa(ml, hl)
        ac1 = kappa.gwet_ac1(ml, hl)
        m_nr = [c for c in dc if c["m_state"] == STATE_NR]
        h_nr = [c for c in dc if c["h_state"] == STATE_NR]
        prec = wilson(sum(c["h_state"] == STATE_NR for c in m_nr), len(m_nr))
        rec = wilson(sum(c["m_state"] == STATE_NR for c in h_nr), len(h_nr))
        types = Counter()
        for c in dc:
            if c["m_label"] == c["h_label"]:
                continue
            t = disagreement_type(c["m_state"], c["m_vals"], c["h_state"], c["h_vals"],
                                  ABSENCE_VALUE.get(d["key"], "none"))
            types[t] += 1
            m, h = c["model"], c["human"]
            disagreements.append({
                "cell_id": c["cell_id"], "system_id": c["system_id"], "dim_id": d["id"],
                "dim_key": d["key"], "stratum": c["stratum"], "type": t,
                "model_label": c["m_label"], "human_label": c["h_label"],
                "model_evidence": m.get("model_evidence", ""), "model_note": m.get("model_note", ""),
                "human_evidence": (h.get("human_evidence_quote") or "").strip(),
                "human_locator": (h.get("human_locator") or "").strip(),
                "human_note": (h.get("human_note") or "").strip(), "adjudication": "",
            })
        row = {
            "dim_id": d["id"], "name": d["name"], "multi": bool(d.get("multi")),
            "n_cells": len(dc), "n_human_unresolved": len(dc) - len(ref), "n_ref": len(ref),
            "n_correct": k_ok, "accuracy": acc, "accuracy_ci_lo": lo, "accuracy_ci_hi": hi,
            "state_confusion_model_by_human": confusion, "state_agreement": st_po,
            "state_kappa": st_k, "n_both_coded": len(both), "value_agreement_given_coded": v_p,
            "value_agreement_ci_lo": v_lo, "value_agreement_ci_hi": v_hi,
            "kappa": k_lab, "observed_agreement": po_lab, "ac1": ac1,
            "silence_precision": prec[0], "silence_precision_ci": [prec[1], prec[2]],
            "n_model_nr": len(m_nr), "silence_recall": rec[0],
            "silence_recall_ci": [rec[1], rec[2]], "n_human_nr": len(h_nr),
            "disagreement_types": dict(types),
        }
        if d.get("multi"):
            pv, pv_po, _ = kappa.multilabel_kappa(ml, hl)
            row["per_value_kappa"] = pv
            row["per_value_agreement"] = pv_po
        per_dim[d["key"]] = row
    ref_all = [c for c in cells if c["h_state"] != STATE_UNRESOLVED]
    k_all = sum(c["correct"] for c in ref_all)
    p_all, lo_all, hi_all = wilson(k_all, len(ref_all))
    _, wacc = pooled_accuracy(cells)
    boot = cluster_bootstrap_accuracy(cells, draws, seed)
    per_stratum = {}
    for st in STRATA:
        sc = [c for c in ref_all if c["stratum"] == st]
        a, lo, hi = wilson(sum(c["correct"] for c in sc), len(sc))
        per_stratum[st] = {"n_ref": len(sc), "n_systems": len({c["system_id"] for c in sc}),
                           "accuracy": a, "accuracy_ci_lo": lo, "accuracy_ci_hi": hi}
    minutes = [float(c["human"]["human_minutes"]) for c in cells
               if (c["human"].get("human_minutes") or "").strip()]
    complete = len(cells) == n_sheet_rows
    return _clean({
        "status": "complete" if complete else "partial",
        "n_sheet_rows": n_sheet_rows, "n_filled": len(cells),
        "n_systems": len({c["system_id"] for c in cells}),
        "pooled": {
            "n_ref": len(ref_all), "n_correct": k_all, "accuracy": p_all,
            "accuracy_wilson_ci_lo": lo_all, "accuracy_wilson_ci_hi": hi_all,
            "weighted_accuracy": wacc, "paper_cluster_bootstrap": boot,
        },
        "per_stratum": per_stratum,
        "per_dimension": per_dim,
        "disagreement_types": dict(Counter(x["type"] for x in disagreements)),
        "human_minutes_total": sum(minutes) if minutes else None,
        "human_minutes_median": sorted(minutes)[len(minutes) // 2] if minutes else None,
        "disagreements": disagreements,
    })


# -------------------------------------------------------------------------------------- outputs


def _f(x, nd: int = 3) -> str:
    return "n/a" if x is None else f"{x:.{nd}f}"


def _md(s: str) -> str:
    return " ".join(str(s or "").split()).replace("|", "\\|")


def results_markdown(res: dict) -> str:
    p = res["pooled"]
    b = p["paper_cluster_bootstrap"]
    lines = [
        "# Human correctness audit: results",
        "",
        ("GENERATED by `scripts/human_audit.py --analyse`; do not edit. Protocol: "
        "`docs/human_audit_protocol.md`."),
        "",
        (f"Status: **{res['status']}** ({res['n_filled']} of {res['n_sheet_rows']} cells filled, "
        f"{res['n_systems']} systems)."),
        "",
        "## Pooled accuracy of the model reading against the human reading",
        "",
        "| measure | value | 95% interval |",
        "|---|---|---|",
        (f"| accuracy ({p['n_correct']}/{p['n_ref']} cells) | {_f(p['accuracy'])} | Wilson "
        f"[{_f(p['accuracy_wilson_ci_lo'])}, {_f(p['accuracy_wilson_ci_hi'])}]; paper-clustered "
        f"bootstrap [{_f(b['accuracy_ci_lo'])}, {_f(b['accuracy_ci_hi'])}] |"),
        (f"| frame-weighted accuracy | {_f(p['weighted_accuracy'])} | paper-clustered bootstrap "
        f"[{_f(b['weighted_accuracy_ci_lo'])}, {_f(b['weighted_accuracy_ci_hi'])}] |"),
        "",
        f"Bootstrap: {b['draws']} draws over {b['n_clusters']} paper clusters, seed {b['seed']}.",
        "",
        "## Per dimension",
        "",
        ("| dim | n ref | accuracy [Wilson 95%] | state agr. | state kappa | value agr. given "
        "both coded (n) | kappa | AC1 | silence precision (n) | silence recall (n) |"),
        "|---|---|---|---|---|---|---|---|---|---|",
    ]
    for key, r in res["per_dimension"].items():
        lines.append(
            f"| {r['dim_id']} {key} | {r['n_ref']} | {_f(r['accuracy'])} "
            f"[{_f(r['accuracy_ci_lo'])}, {_f(r['accuracy_ci_hi'])}] | "
            f"{_f(r['state_agreement'])} | {_f(r['state_kappa'])} | "
            f"{_f(r['value_agreement_given_coded'])} ({r['n_both_coded']}) | {_f(r['kappa'])} | "
            f"{_f(r['ac1'])} | {_f(r['silence_precision'])} ({r['n_model_nr']}) | "
            f"{_f(r['silence_recall'])} ({r['n_human_nr']}) |")
    multi = [(k, r) for k, r in res["per_dimension"].items() if "per_value_kappa" in r]
    for k, r in multi:
        lines += ["", (f"{r['dim_id']} {k} (multi-valued): per-value kappa "
                      f"{_f(r['per_value_kappa'])}, per-value agreement "
                      f"{_f(r['per_value_agreement'])} (diagnostic only).")]
    lines += ["", "## State confusion (rows: model, columns: human)", ""]
    for key, r in res["per_dimension"].items():
        cm = r["state_confusion_model_by_human"]
        lines += [f"**{r['dim_id']} {key}**", "", "| model \\ human | " + " | ".join(STATES) + " |",
                  "|---|---|---|---|"]
        for ms in STATES:
            lines.append(f"| {ms} | " + " | ".join(str(cm[ms][hs]) for hs in STATES) + " |")
        lines.append("")
    lines += ["## Per stratum (descriptive)", "", "| stratum | systems | n ref | accuracy [Wilson] |",
              "|---|---|---|---|"]
    for st, r in res["per_stratum"].items():
        lines.append(f"| {st} | {r['n_systems']} | {r['n_ref']} | {_f(r['accuracy'])} "
                     f"[{_f(r['accuracy_ci_lo'])}, {_f(r['accuracy_ci_hi'])}] |")
    types = ", ".join(f"{k} {v}" for k, v in sorted(res["disagreement_types"].items())) or "none"
    lines += ["", f"## Disagreements ({len(res['disagreements'])}: {types})", "",
              ("The `adjudication` column is filled after unblinding (protocol 4.4) and never "
              "changes the accuracy above."), "",
              ("| cell | type | model | human | model evidence (note) | human evidence (locator; "
              "note) | adjudication |"), "|---|---|---|---|---|---|---|"]
    for x in res["disagreements"]:
        me = _md(x["model_evidence"]) + (f" ({_md(x['model_note'])})" if x["model_note"] else "")
        he = _md(x["human_evidence"]) + f" ({_md(x['human_locator'])}"
        he += f"; {_md(x['human_note'])})" if x["human_note"] else ")"
        lines.append(f"| {x['cell_id']} | {x['type']} | {_md(x['model_label'])} | "
                     f"{_md(x['human_label'])} | {me} | {he} | {x['adjudication']} |")
    if res.get("human_minutes_total") is not None:
        lines += ["", (f"Human time: {res['human_minutes_total']:.0f} minutes in total, median "
                      f"{res['human_minutes_median']:.1f} per cell.")]
    return "\n".join(lines) + "\n"


def _tex(x) -> str:
    return "--" if x is None else f"{x:.2f}"


def results_tex(res: dict, strata: dict[str, int]) -> str:
    if res["status"] != "complete":
        raise AuditError("refusing to write the LaTeX fragment from a partial sheet")
    p = res["pooled"]
    b = p["paper_cluster_bootstrap"]
    st = " / ".join(f"{k} {v}" for k, v in strata.items())
    out = [
        "% Human correctness audit",
        "% GENERATED by scripts/human_audit.py --analyse - do not edit; re-run the script instead.",
        "% protocol: docs/human_audit_protocol.md; data: data/audit/results.json",
        "\\begin{table}",
        "\\centering",
        (f"\\caption{{Blinded human audit of the released model coding: {res['n_systems']} systems "
        f"({st}) not in the reliability sample, 7 dimensions, {p['n_ref']} cells with a human "
        "reference. Accuracy is exact agreement of state and value set with the human reading, "
        "with a Wilson 95\\% interval; pooled interval from a paper-clustered bootstrap. "
        "\\emph{Silence prec.}: share of model \\emph{not reported} cells the human also found "
        "not reported.}"),
        "\\label{tab:human-audit}",
        "\\footnotesize",
        "\\begin{tabular}{@{}lrlrrrrr@{}}",
        "\\toprule",
        ("Dimension & $n$ & Accuracy [95\\% CI] & State agr. & Value agr.$^a$ & $\\kappa$ & AC1 "
        "& Silence prec. \\\\"),
        "\\midrule",
        "% BEGIN ROWS",
    ]
    for key, r in res["per_dimension"].items():
        name = key.replace("_", " ")
        out.append(
            f"{r['dim_id']} {name} & {r['n_ref']} & {_tex(r['accuracy'])} "
            f"[{_tex(r['accuracy_ci_lo'])}, {_tex(r['accuracy_ci_hi'])}] & "
            f"{_tex(r['state_agreement'])} & {_tex(r['value_agreement_given_coded'])} "
            f"({r['n_both_coded']}) & {_tex(r['kappa'])} & {_tex(r['ac1'])} & "
            f"{_tex(r['silence_precision'])} \\\\")
    out += [
        "\\midrule",
        (f"Pooled & {p['n_ref']} & {_tex(p['accuracy'])} [{_tex(b['accuracy_ci_lo'])}, "
        f"{_tex(b['accuracy_ci_hi'])}] & & & & & \\\\"),
        (f"Pooled, frame-weighted & & {_tex(p['weighted_accuracy'])} "
        f"[{_tex(b['weighted_accuracy_ci_lo'])}, {_tex(b['weighted_accuracy_ci_hi'])}] "
        "& & & & & \\\\"),
        "% END ROWS",
        "\\bottomrule",
        "\\end{tabular}",
        "",
        ("\\raggedright\\footnotesize $^a$ Exact value-set agreement among cells both readings coded "
        "($n$ in parentheses)."),
        "\\end{table}",
    ]
    return "\n".join(out) + "\n"


def pick_sheet(out_dir: Path, explicit: str | None) -> Path:
    if explicit:
        return Path(explicit)
    cands = [out_dir / "human_sheet.csv", out_dir / "human_sheet.xlsx"]
    used = [p for p in cands if p.exists() and any(filled(r) for r in read_sheet(p))]
    if len(used) > 1:
        raise AuditError("both human_sheet.csv and human_sheet.xlsx carry entries; name the one "
                         "to analyse with --sheet")
    if used:
        return used[0]
    return cands[0]


def run_analysis(root: Path, out_dir: Path, sheet_path: Path, tex_path: Path | None,
                 allow_partial: bool = False, draws: int = BOOTSTRAP_DRAWS) -> dict:
    if not sheet_path.exists():
        raise AuditError(f"no sheet at {sheet_path}; run --build first")
    sheet = read_sheet(sheet_path)
    model = read_csv(out_dir / "model_answers.csv")
    sample = read_csv(out_dir / "sample.csv")
    dims = audit_dims(read_json(root / "schema" / "dimensions.json"))
    human = validate(sheet, model, dims, allow_partial)
    res = analyse(human, model, sample, dims, len(sheet), draws=draws)
    res["sheet"] = str(sheet_path)
    (out_dir / "results.json").write_text(json.dumps(res, indent=1, ensure_ascii=False) + "\n",
                                          encoding="utf-8")
    (out_dir / "results.md").write_text(results_markdown(res), encoding="utf-8")
    if tex_path is not None and res["status"] == "complete":
        strata = dict(sorted(Counter(r["stratum"] for r in sample).items(),
                             key=lambda kv: STRATA.index(kv[0])))
        tex_path.parent.mkdir(parents=True, exist_ok=True)
        tex_path.write_text(results_tex(res, strata), encoding="utf-8")
    return res


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument("--build", action="store_true", help="draw the sample and write the sheets")
    g.add_argument("--analyse", "--analyze", action="store_true", help="score a filled sheet")
    ap.add_argument("--root", default=str(ROOT))
    ap.add_argument("--out-dir", default=None, help="default <root>/data/audit")
    ap.add_argument("--sheet", default=None, help="filled sheet (.csv or .xlsx)")
    ap.add_argument("--tex", default=None, help="default <root>/paper/tables/human_audit.tex")
    ap.add_argument("--no-tex", action="store_true")
    ap.add_argument("--allow-partial", action="store_true")
    ap.add_argument("--draws", type=int, default=BOOTSTRAP_DRAWS)
    ap.add_argument("--force", action="store_true", help="overwrite a sheet with human entries")
    a = ap.parse_args(argv)
    root = Path(a.root)
    out_dir = Path(a.out_dir) if a.out_dir else root / "data" / "audit"
    try:
        if a.build:
            info = build(root, out_dir, force=a.force)
            print(json.dumps(info, indent=1))
            return 0
        tex = None if a.no_tex else Path(a.tex) if a.tex else root / "paper" / "tables" / \
            "human_audit.tex"
        res = run_analysis(root, out_dir, pick_sheet(out_dir, a.sheet), tex, a.allow_partial,
                           a.draws)
        p = res["pooled"]
        print(f"{res['status']}: {res['n_filled']} cells; pooled accuracy {_f(p['accuracy'])} "
              f"(weighted {_f(p['weighted_accuracy'])}); {len(res['disagreements'])} "
              "disagreements")
        if res["status"] != "complete":
            print("partial sheet: LaTeX fragment NOT written")
        return 0
    except AuditError as e:
        print(f"error: {e}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
